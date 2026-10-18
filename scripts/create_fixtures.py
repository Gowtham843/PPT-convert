#!/usr/bin/env python3
"""
Create Phase 1 test PPTX fixtures under test-files/ (01–06).

Methods:
  - python-pptx for text, shapes, images, video, audio
  - Direct OOXML injection for timing (animations) and transitions

Animation/transition XML is constructed to match ECMA-376-like shapes.
Not validated in Microsoft PowerPoint UI on this machine.
"""

from __future__ import annotations

import subprocess
import struct
import wave
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "test-files"
MEDIA = OUT / "_assets"
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def ensure_media():
    MEDIA.mkdir(parents=True, exist_ok=True)
    png = MEDIA / "sample.png"
    if not png.exists():
        try:
            from PIL import Image

            Image.new("RGB", (120, 90), (200, 50, 50)).save(png)
        except ImportError:
            subprocess.run(
                ["ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=red:s=120x90", str(png)],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
    wav = MEDIA / "sample.wav"
    if not wav.exists():
        with wave.open(str(wav), "w") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(22050)
            w.writeframes(struct.pack("<h", 0) * 22050)
    mp4 = MEDIA / "sample.mp4"
    if not mp4.exists():
        subprocess.run(
            [
                "ffmpeg",
                "-y",
                "-f",
                "lavfi",
                "-i",
                "color=c=blue:s=320x240:d=2",
                "-f",
                "lavfi",
                "-i",
                "sine=f=440:d=2",
                "-shortest",
                str(mp4),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    return png, wav, mp4


def save(prs: Presentation, name: str) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{name}.pptx"
    prs.save(str(path))
    return path


def inject_slide_xml(pptx_path: Path, slide_idx: int, mutator):
    """Mutate slide XML inside pptx zip (slide_idx is 1-based)."""
    slide_name = f"ppt/slides/slide{slide_idx}.xml"
    with zipfile.ZipFile(pptx_path, "r") as zin:
        data = zin.read(slide_name)
        root = etree.fromstring(data)
        mutator(root)
        new_data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        tmp = pptx_path.with_suffix(".tmp.pptx")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == slide_name:
                    zout.writestr(item, new_data)
                else:
                    zout.writestr(item, zin.read(item.filename))
    tmp.replace(pptx_path)


def add_transition(root, kind: str):
    """Add p:transition child to slide root."""
    for child in list(root):
        if child.tag == qn("p:transition"):
            root.remove(child)
    trans = etree.Element(qn("p:transition"), spd="med")
    if kind == "fade":
        etree.SubElement(trans, qn("p:fade"))
    elif kind == "push":
        push = etree.SubElement(trans, qn("p:push"), dir="l")
    elif kind == "wipe":
        etree.SubElement(trans, qn("p:wipe"), dir="l")
    root.append(trans)


def add_fade_in_animation(root, spid: str):
    """Inject minimal timing tree for fade-in on shape spid."""
    for child in list(root):
        if child.tag == qn("p:timing"):
            root.remove(child)
    timing = etree.Element(qn("p:timing"))
    tn_lst = etree.SubElement(timing, qn("p:tnLst"))
    par = etree.SubElement(tn_lst, qn("p:par"))
    ctn = etree.SubElement(par, qn("p:cTn"), id="1", dur="indefinite", restart="never", nodeType="tmRoot")
    child_tn = etree.SubElement(ctn, qn("p:childTnLst"))
    seq = etree.SubElement(child_tn, qn("p:seq"), concurrent="1", nextAc="seek")
    seq_ctn = etree.SubElement(seq, qn("p:cTn"), id="2", dur="indefinite", nodeType="mainSeq")
    seq_child = etree.SubElement(seq_ctn, qn("p:childTnLst"))
    click_par = etree.SubElement(seq_child, qn("p:par"))
    click_ctn = etree.SubElement(
        click_par, qn("p:cTn"), id="3", fill="hold", presetID="1", presetClass="entr", nodeType="clickEffect"
    )
    click_child = etree.SubElement(click_ctn, qn("p:childTnLst"))
    inner_par = etree.SubElement(click_child, qn("p:par"))
    inner_ctn = etree.SubElement(inner_par, qn("p:cTn"), id="4", fill="hold", dur="500")
    inner_child = etree.SubElement(inner_ctn, qn("p:childTnLst"))
    ae = etree.SubElement(inner_child, qn("p:animEffect"), transition="in", filter="fade")
    ae_ctn = etree.SubElement(ae, qn("p:cTn"), id="5", dur="500")
    tgt = etree.SubElement(ae, qn("p:tgtEl"))
    etree.SubElement(tgt, qn("p:spTgt"), spid=spid)
    root.append(timing)


def shape_spid(slide, shape_index: int = 0) -> str:
    shape = slide.shapes[shape_index]
    el = shape._element
    cnv = el.find(".//p:cNvPr", namespaces={"p": NS["p"]})
    return cnv.get("id") if cnv is not None else "2"


def fixture_static():
    png, _, _ = ensure_media()
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    box = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8), Inches(1))
    box.text_frame.text = "Static fixture — text, shapes, image"
    box.text_frame.paragraphs[0].font.size = Pt(28)
    rect = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1.5))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(68, 114, 196)
    rect.text_frame.text = "Shape"
    s1.shapes.add_picture(str(png), Inches(5), Inches(2), width=Inches(2.5))

    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "Table slide (PNG fallback expected)"
    rows, cols = 3, 3
    table_shape = s2.shapes.add_table(rows, cols, Inches(1), Inches(1.5), Inches(7), Inches(3))
    table = table_shape.table
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = f"R{r+1}C{c+1}"

    path = save(prs, "01-static")
    print("✓", path.name)


def fixture_video():
    _, _, mp4 = ensure_media()
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "Video fixture"
    s.shapes.add_movie(str(mp4), Inches(1.5), Inches(1.5), Inches(6), Inches(4))
    print("✓", save(prs, "02-video").name)


def fixture_audio():
    _, wav, _ = ensure_media()
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "Audio fixture"
    s.shapes.add_movie(str(wav), Inches(2), Inches(3), Inches(5), Inches(0.6))
    print("✓", save(prs, "03-audio").name)


def fixture_animation():
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "Animation fixture (fade in box)"
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(2.5), Inches(4), Inches(1.5))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(237, 125, 49)
    box.text_frame.text = "Animate me"
    spid = shape_spid(s, 1)
    path = save(prs, "04-animation")
    inject_slide_xml(path, 1, lambda root: add_fade_in_animation(root, spid))
    print("✓", path.name, f"(spid={spid})")


def fixture_transition():
    prs = Presentation()
    blank = prs.slide_layouts[6]
    kinds = [("fade", "Fade transition"), ("push", "Push transition"), ("wipe", "Wipe transition")]
    for kind, title in kinds:
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    path = save(prs, "05-transition")
    for i, (kind, _) in enumerate(kinds, start=1):
        inject_slide_xml(path, i, lambda root, k=kind: add_transition(root, k))
    print("✓", path.name)


def fixture_combined():
    png, wav, mp4 = ensure_media()
    prs = Presentation()
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "Combined — static + media"
    s1.shapes.add_picture(str(png), Inches(0.5), Inches(1.2), width=Inches(2))
    s1.shapes.add_movie(str(mp4), Inches(3), Inches(1.2), Inches(5), Inches(3))

    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)).text_frame.text = "Combined — animation"
    anim_box = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(2.5), Inches(3), Inches(2))
    anim_box.fill.solid()
    anim_box.fill.fore_color.rgb = RGBColor(112, 173, 71)
    spid = shape_spid(s2, 1)

    s3 = prs.slides.add_slide(blank)
    s3.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1)).text_frame.text = "Combined — fade transition"
    s3.shapes.add_movie(str(wav), Inches(2), Inches(4), Inches(5), Inches(0.6))

    path = save(prs, "06-combined")
    inject_slide_xml(path, 2, lambda root: add_fade_in_animation(root, spid))
    inject_slide_xml(path, 3, lambda root: add_transition(root, "fade"))
    print("✓", path.name)


def main():
    ensure_media()
    OUT.mkdir(parents=True, exist_ok=True)
    fixture_static()
    fixture_video()
    fixture_audio()
    fixture_animation()
    fixture_transition()
    fixture_combined()
    print(f"\nFixtures written to {OUT}")


if __name__ == "__main__":
    main()
